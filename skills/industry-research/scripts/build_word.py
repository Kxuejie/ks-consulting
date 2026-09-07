"""Build an editable research report from the JSON contract in references/output.md."""
import argparse
import json
from pathlib import Path
from urllib.parse import urlparse


def validate(data):
    for key in ('title', 'subtitle', 'summary'):
        if not isinstance(data.get(key), str) or not data[key].strip():
            raise ValueError(f'{key} must be nonempty text')
    if not isinstance(data.get('sections'), list) or not data['sections']:
        raise ValueError('sections must be a nonempty list')
    for section in data['sections']:
        if not isinstance(section.get('title'), str) or not section['title'].strip():
            raise ValueError('section title must be nonempty text')
        if not isinstance(section.get('blocks'), list) or not section['blocks']:
            raise ValueError('section blocks must be a nonempty list')
        for block in section['blocks']:
            kind = block.get('type')
            if kind in ('paragraph', 'heading'):
                if not isinstance(block.get('text'), str):
                    raise ValueError('text must be a string')
                if kind == 'heading' and block.get('level', 2) not in (2, 3, 4):
                    raise ValueError('block heading level must be 2, 3 or 4')
            elif kind == 'bullets':
                if not isinstance(block.get('items'), list) or not all(isinstance(x, str) for x in block['items']):
                    raise ValueError('bullet items must be text')
            elif kind == 'links':
                if not isinstance(block.get('items'), list):
                    raise ValueError('links items must be a list')
                for item in block['items']:
                    if not isinstance(item.get('text'), str) or not isinstance(item.get('url'), str):
                        raise ValueError('link text and url must be strings')
                    parsed = urlparse(item['url'])
                    if parsed.scheme not in ('https', 'http') or not parsed.netloc:
                        raise ValueError('source links must be absolute HTTP(S) URLs')
            elif kind == 'table':
                headers, rows = block.get('headers'), block.get('rows')
                if not isinstance(headers, list) or not headers or not all(isinstance(x, str) for x in headers):
                    raise ValueError('table headers must be nonempty text list')
                if not isinstance(rows, list) or not all(isinstance(r, list) and len(r) == len(headers) and all(isinstance(x, str) for x in r) for r in rows):
                    raise ValueError('table rows must match header count and contain strings')
                widths = block.get('widths', [1] * len(headers))
                if len(widths) != len(headers) or not all(type(x) in (int, float) and 0 < x < 1000 for x in widths):
                    raise ValueError('widths must contain positive relative widths')
            else:
                raise ValueError(f'unsupported block type: {kind}')


def build(data, output):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.opc.constants import RELATIONSHIP_TYPE

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(.8)
    sec.left_margin = sec.right_margin = Inches(.85)
    for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Heading 4', 'List Bullet'):
        style = doc.styles[name]
        style.font.name = 'Calibri'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Microsoft YaHei')
        style.paragraph_format.space_after = Pt(8)
    doc.styles['Normal'].font.size = Pt(11)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.25
    for name, size in (('Title', 26), ('Heading 1', 18), ('Heading 2', 14), ('Heading 3', 12), ('Heading 4', 11)):
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].paragraph_format.keep_with_next = True
    doc.add_paragraph(data['title'], 'Title')
    doc.add_paragraph(data['subtitle'], 'Subtitle')
    doc.add_paragraph(data['summary'])
    for section in data['sections']:
        doc.add_heading(section['title'], level=1)
        for block in section['blocks']:
            kind = block['type']
            if kind == 'heading':
                doc.add_heading(block['text'], level=block.get('level', 2))
            elif kind == 'paragraph':
                doc.add_paragraph(block['text'])
            elif kind == 'bullets':
                for item in block['items']:
                    doc.add_paragraph(item, 'List Bullet')
            elif kind == 'links':
                for item in block['items']:
                    p = doc.add_paragraph()
                    link = OxmlElement('w:hyperlink')
                    link.set(qn('r:id'), p.part.relate_to(item['url'], RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
                    run = OxmlElement('w:r')
                    props = OxmlElement('w:rPr')
                    color = OxmlElement('w:color'); color.set(qn('w:val'), '416F9F'); props.append(color)
                    run.append(props)
                    text = OxmlElement('w:t'); text.text = item['text']; run.append(text)
                    link.append(run); p._p.append(link)
            elif kind == 'table':
                table = doc.add_table(rows=1, cols=len(block['headers']))
                table.autofit = False
                ratios = block.get('widths', [1] * len(block['headers']))
                widths = [Inches(6.8 * x / sum(ratios)) for x in ratios]
                for col, width in zip(table.columns, widths):
                    col.width = width
                for cell, value in zip(table.rows[0].cells, block['headers']):
                    cell.text = value
                repeat = OxmlElement('w:tblHeader')
                table.rows[0]._tr.get_or_add_trPr().append(repeat)
                for values in block['rows']:
                    for cell, value in zip(table.add_row().cells, values):
                        cell.text = value
                for index, row in enumerate(table.rows):
                    for cell, width in zip(row.cells, widths):
                        cell.width = width
                        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        props = cell._tc.get_or_add_tcPr()
                        shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E5EDF5' if index == 0 else ('FFFFFF' if index % 2 else 'F7F8FA')); props.append(shade)
                        borders = OxmlElement('w:tcBorders')
                        margins = OxmlElement('w:tcMar')
                        for edge in ('top', 'left', 'bottom', 'right'):
                            border = OxmlElement(f'w:{edge}')
                            for key, value in (('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')):
                                border.set(qn(f'w:{key}'), value)
                            borders.append(border)
                            margin = OxmlElement(f'w:{edge}'); margin.set(qn('w:w'), '100'); margin.set(qn('w:type'), 'dxa'); margins.append(margin)
                        props.append(borders); props.append(margins)
                        for p in cell.paragraphs:
                            p.paragraph_format.space_after = Pt(4)
                            p.paragraph_format.space_before = Pt(4)
                            for run in p.runs:
                                run.font.size = Pt(10.5)
                                run.bold = index == 0
                doc.add_paragraph()
    footer = sec.footer.paragraphs[0]
    footer.alignment = 2
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    doc.core_properties.title = data['title']
    doc.core_properties.subject = data['subtitle']
    doc.save(output)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input', type=Path)
    parser.add_argument('output', type=Path)
    parser.add_argument('--overwrite', action='store_true')
    args = parser.parse_args()
    if args.output.suffix.lower() != '.docx':
        parser.error('output must use .docx')
    if args.output.exists() and not args.overwrite:
        parser.error('output already exists; choose another name or use --overwrite')
    data = json.loads(args.input.read_text(encoding='utf-8-sig'))
    validate(data)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    build(data, args.output)
    print(args.output.resolve())


if __name__ == '__main__':
    main()
